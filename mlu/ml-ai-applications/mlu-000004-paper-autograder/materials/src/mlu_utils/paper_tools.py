"""Paper autograding for the Paper Autograding Lab.

Keeps paper loading, rubric handling, roster reconciliation, the grounded
grading prompt + model, and report rendering OUT of the notebook so the
notebook stays use-led rather than a wall of code (same convention as
mlu_utils/embeddings.py and the participation lab's transcript_tools.py).

The notebook does:

    from mlu_utils.paper_tools import (
        load_rubric, load_papers, reconcile_roster,
        PaperGrader, build_report, write_student_feedback,
    )

    rubric      = load_rubric("data/sample_rubric.py")
    submissions = load_papers("data/sample_papers")
    matched, missing, off_roster = reconcile_roster(submissions, roster_names)

    grader = PaperGrader(retriever, bedrock_runtime, rubric)
    graded = grader.grade_all(matched, missing)
    report_md, csv_rows = build_report(graded, rubric)

WHAT THIS IS: a DRAFT paper grade the instructor reviews. Every judgment that
touches "did the student use the course concepts / get the material right" is
grounded ONLY in the uploaded course material via the retriever — the same RAG
stack the participation and curriculum labs use. The overall score is computed
in CODE from the per-criterion tiers and the rubric weights, so the arithmetic
is never the model's job. The grade never reaches a gradebook without a human
looking at it. See Parts 0 and 7 of the notebook.

WHAT THIS IS NOT: a plagiarism or AI-writing detector, and not an academic-
integrity tool. It reads a paper and scores it against a rubric. Do not use its
output to accuse a student of misconduct.
"""
import csv
import json
import os
import re
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

from langchain_core.output_parsers import StrOutputParser
from langchain.prompts import PromptTemplate

# NOTE: the chat model is NOT imported at module load. Bedrock (ChatBedrockConverse)
# is imported lazily only when the grader has to build it itself, so a local
# Ollama-only install without langchain-aws can still `import paper_tools`.


# --------------------------------------------------------------------------- #
# Paper loading  (.pdf via pypdf, .docx via python-docx, .txt/.md via stdlib) #
# --------------------------------------------------------------------------- #

@dataclass
class Submission:
    """One student's paper. `student_name` is inferred from the filename
    (`Aisha_Rahman.pdf` -> "Aisha Rahman") unless a roster overrides it."""
    student_name: str
    path: str
    text: str

    @property
    def words(self) -> int:
        return len(self.text.split())


_SUPPORTED_EXT = (".pdf", ".docx", ".txt", ".md")


def _name_from_filename(path: str) -> str:
    """Turn a submission filename into a student display name. Strips the
    extension and any leading numbering, and treats '_' / '-' as spaces —
    'Aisha_Rahman.pdf', '03-rahman-aisha.docx' both become readable names."""
    stem = os.path.splitext(os.path.basename(path))[0]
    stem = re.sub(r"^\s*\d+[\s._-]+", "", stem)          # drop a leading "03_" index
    stem = re.sub(r"[._-]+", " ", stem).strip()
    return re.sub(r"\s+", " ", stem)


def _read_pdf(path: str) -> str:
    from pypdf import PdfReader
    reader = PdfReader(path)
    return "\n".join((page.extract_text() or "") for page in reader.pages).strip()


def _read_docx(path: str) -> str:
    try:
        from docx import Document
    except ImportError as e:                              # pragma: no cover
        raise ImportError(
            "Reading .docx papers needs python-docx. `pip install python-docx`, "
            "or ask students to submit PDF/TXT."
        ) from e
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    # tables (rubric-style appendices, data tables) hold real content too
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts).strip()


def _read_text(path: str) -> str:
    with open(path, "r", encoding="utf-8-sig") as f:
        return f.read().strip()


def load_paper(path: str, student_name: Optional[str] = None) -> Submission:
    """Load one paper from .pdf, .docx, .txt, or .md into a Submission."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        text = _read_pdf(path)
    elif ext == ".docx":
        text = _read_docx(path)
    elif ext in (".txt", ".md"):
        text = _read_text(path)
    else:
        raise ValueError(f"Unsupported paper format '{ext}'. Use one of {_SUPPORTED_EXT}.")
    return Submission(
        student_name=student_name or _name_from_filename(path),
        path=path,
        text=text,
    )


def load_papers(folder: str) -> List[Submission]:
    """Load every supported paper in a folder — one Submission per file.

    Files that fail to read (a corrupt PDF, a scanned image with no text layer)
    are returned with empty text rather than aborting the whole batch; the
    grader flags an empty paper as 'unreadable' for the instructor to chase."""
    submissions: List[Submission] = []
    for name in sorted(os.listdir(folder)):
        if name.startswith(".") or os.path.splitext(name)[1].lower() not in _SUPPORTED_EXT:
            continue
        path = os.path.join(folder, name)
        if not os.path.isfile(path):
            continue
        try:
            submissions.append(load_paper(path))
        except Exception:                                 # noqa: BLE001 - degrade, don't crash the batch
            submissions.append(Submission(student_name=_name_from_filename(path), path=path, text=""))
    return submissions


# --------------------------------------------------------------------------- #
# Rubric  (structured criteria + weights + level descriptors)                 #
# --------------------------------------------------------------------------- #

@dataclass
class Criterion:
    name: str
    weight: float                      # relative weight; normalized across criteria
    description: str = ""
    levels: Dict[int, str] = field(default_factory=dict)   # tier 0-4 -> descriptor

    def rubric_block(self) -> str:
        """Render this criterion for the grading prompt."""
        lines = [f"CRITERION: {self.name}  (worth {self.weight:g}% of the grade)"]
        if self.description:
            lines.append(f"  What it measures: {self.description}")
        for tier in sorted(self.levels, reverse=True):
            lines.append(f"  {tier} — {self.levels[tier]}")
        return "\n".join(lines)


@dataclass
class Rubric:
    criteria: List[Criterion]
    grade_bands: List[Tuple[float, str]] = field(default_factory=list)  # (min_pct, letter), desc
    max_tier: int = 4

    def normalized_weights(self) -> Dict[str, float]:
        """Weights as fractions summing to 1.0, so a rubric whose weights sum to
        90 or 103 still produces a clean 0-100 overall instead of silently
        under- or over-counting."""
        total = sum(c.weight for c in self.criteria) or 1.0
        return {c.name: c.weight / total for c in self.criteria}

    def overall_percent(self, tiers: Dict[str, Optional[int]]) -> Optional[float]:
        """Weighted overall as a 0-100 percentage, computed in code from the
        per-criterion tiers. A criterion the model failed to score (None) is
        dropped and the remaining weights are renormalized, so one bad criterion
        does not zero the paper — it just widens the confidence gap."""
        w = self.normalized_weights()
        num = 0.0
        denom = 0.0
        for c in self.criteria:
            t = tiers.get(c.name)
            if t is None:
                continue
            num += w[c.name] * (t / self.max_tier)
            denom += w[c.name]
        if denom == 0:
            return None
        return round(100.0 * num / denom, 1)

    def letter(self, percent: Optional[float]) -> str:
        if percent is None or not self.grade_bands:
            return "—"
        for lo, letter in sorted(self.grade_bands, reverse=True):
            if percent >= lo:
                return letter
        return self.grade_bands[-1][1] if self.grade_bands else "—"


DEFAULT_GRADE_BANDS: List[Tuple[float, str]] = [
    (93, "A"), (90, "A-"), (87, "B+"), (83, "B"), (80, "B-"),
    (77, "C+"), (73, "C"), (70, "C-"), (60, "D"), (0, "F"),
]


def load_rubric(path: str) -> Rubric:
    """Load a structured rubric from a Python file exposing `RUBRIC` (a dict) or
    a `.json` file with the same shape:

        RUBRIC = {
          "max_tier": 4,
          "grade_bands": [[93,"A"],[90,"A-"], ...],   # optional; sensible default
          "criteria": [
            {"name": "...", "weight": 30, "description": "...",
             "levels": {4: "...", 3: "...", 2: "...", 1: "...", 0: "..."}},
            ...
          ],
        }
    """
    ext = os.path.splitext(path)[1].lower()
    if ext == ".json":
        with open(path, "r", encoding="utf-8-sig") as f:
            data = json.load(f)
    elif ext == ".py":
        ns: Dict[str, object] = {}
        with open(path, "r", encoding="utf-8-sig") as f:
            exec(compile(f.read(), path, "exec"), ns)     # trusted, faculty-authored file
        data = ns.get("RUBRIC")
        if not isinstance(data, dict):
            raise ValueError(f"{path} must define a dict named RUBRIC.")
    else:
        raise ValueError("Rubric must be a .py (defining RUBRIC) or .json file.")

    criteria = [
        Criterion(
            name=str(c["name"]),
            weight=float(c.get("weight", 1)),
            description=str(c.get("description", "")),
            levels={int(k): str(v) for k, v in (c.get("levels") or {}).items()},
        )
        for c in data.get("criteria", [])
    ]
    if not criteria:
        raise ValueError("Rubric has no criteria.")
    bands = [(float(lo), str(lab)) for lo, lab in data.get("grade_bands", DEFAULT_GRADE_BANDS)]
    return Rubric(criteria=criteria, grade_bands=bands, max_tier=int(data.get("max_tier", 4)))


# --------------------------------------------------------------------------- #
# Roster reconciliation                                                       #
# --------------------------------------------------------------------------- #

def load_roster(path: str) -> List[str]:
    """Read a roster CSV and return student names. Looks for a
    `student_name`/`name` column; falls back to the first column."""
    names: List[str] = []
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        cols = reader.fieldnames or []
        key = next((c for c in cols if c and c.strip().lower() in ("student_name", "name", "student")), None)
        if key is None and cols:
            key = cols[0]
        for row in reader:
            val = (row.get(key) or "").strip() if key else ""
            if val:
                names.append(val)
    return names


def _norm(name: str) -> str:
    return re.sub(r"\s+", " ", name.strip().lower())


def _tokens(name: str) -> List[str]:
    return [t for t in re.split(r"\s+", _norm(name)) if t]


def _names_match(a: str, b: str) -> bool:
    """Lenient match: exact normalized equality, or same last name AND first
    names agree on at least an initial. Handles 'Aisha Rahman' vs 'Rahman,
    Aisha' and 'J. Chen' vs 'James Chen' without matching unrelated people.
    (Same matcher as the participation lab — a filename-derived name is treated
    exactly like a transcript-derived one.)"""
    if _norm(a) == _norm(b):
        return True
    ta, tb = _tokens(a.replace(",", " ")), _tokens(b.replace(",", " "))
    if len(ta) >= 2 and len(tb) >= 2:
        if ta[-1] == tb[-1] and ta[0][:1] == tb[0][:1]:
            return True
        if ta[0] == tb[-1] and ta[-1][:1] == tb[0][:1]:
            return True
    return False


def reconcile_roster(
    submissions: List[Submission],
    roster_names: Optional[List[str]],
) -> Tuple[List[Tuple[str, Submission]], List[str], List[Submission]]:
    """Split submissions against a roster.

    Returns (matched, missing, off_roster):
      - matched    : [(roster_name, Submission)] students who turned a paper in
      - missing    : [roster_name] enrolled but no paper found -> flagged 'missing'
      - off_roster : [Submission] a paper whose name matches nobody on the
                     roster (misnamed file, wrong section) -> flagged, NOT graded

    With no roster, every submission is 'matched' under its filename-derived
    name and the other two lists are empty (a non-submitter is then invisible —
    exactly the participation lab's silent-student tradeoff)."""
    if not roster_names:
        return [(s.student_name, s) for s in submissions], [], []

    matched: List[Tuple[str, Submission]] = []
    missing: List[str] = []
    used = set()

    for roster_name in roster_names:
        hit = next(
            (s for s in submissions if id(s) not in used and _names_match(roster_name, s.student_name)),
            None,
        )
        if hit is not None:
            matched.append((roster_name, hit))
            used.add(id(hit))
        else:
            missing.append(roster_name)

    off_roster = [s for s in submissions if id(s) not in used]
    return matched, missing, off_roster


# --------------------------------------------------------------------------- #
# Grading                                                                     #
# --------------------------------------------------------------------------- #

GRADING_TEMPLATE = """You are helping a faculty member draft a grade for ONE student's paper,
against a rubric the faculty member wrote. Your output is a DRAFT the instructor
will review and adjust — it is NOT a final grade, and it is NOT an academic-
integrity or plagiarism judgment.

Whenever you judge whether the student used course concepts correctly or got the
material right, judge ONLY against the COURSE MATERIAL excerpts below. Do not use
outside knowledge of the subject to decide what is correct or what "should" have
been covered. If the paper goes beyond what the course material shows, treat that
as neutral, not wrong.

COURSE MATERIAL (this is the ground truth for what was taught):
{course_context}

THE RUBRIC — score EACH criterion on its own 0-{max_tier} scale:
{rubric_text}

STUDENT: {student_name}  (paper is about {num_words} words)

THE PAPER (verbatim):
{paper_text}

Score every criterion. For each one:
- Pick the tier whose descriptor best fits, judged against the rubric and (where
  relevance/correctness matters) the course material above.
- Do NOT reward length or confident tone. A long paper that misuses the concepts
  scores LOWER on the concept criteria than a short paper that uses them well.
- Pull 1-2 SHORT verbatim quotes from the paper as evidence for the tier.
- Write a one-sentence rationale for the professor (why this tier, grounded in
  the rubric and the material).
- Write a one-sentence note TO THE STUDENT: specific, actionable, and kind —
  what they did and the single most useful thing to improve.
- If the paper is thin or ambiguous for this criterion, say so and lower your
  confidence rather than guessing.

Then write two short paper-level notes: one FOR THE PROFESSOR summarizing the
grade, and one FOR THE STUDENT (2-3 sentences, encouraging and concrete).

Respond with ONLY a JSON object, no prose before or after, in this exact shape:
{{
  "criteria": [
    {{
      "name": "<exact criterion name from the rubric>",
      "tier": <integer 0-{max_tier}>,
      "evidence_quotes": ["<short verbatim quote>", "<another>"],
      "rationale": "<one sentence for the professor>",
      "student_note": "<one actionable sentence for the student>"
    }}
    // ... one object per criterion, same names as the rubric
  ],
  "professor_summary": "<1-2 sentences: the shape of this grade and anything to double-check>",
  "student_summary": "<2-3 sentences of feedback addressed to the student>",
  "confidence": "<high|medium|low>"
}}"""


@dataclass
class CriterionScore:
    name: str
    tier: Optional[int] = None
    evidence_quotes: List[str] = field(default_factory=list)
    rationale: str = ""
    student_note: str = ""


@dataclass
class PaperGrade:
    name: str
    status: str                        # "graded" | "missing" | "unreadable" | "parse_error"
    criteria: List[CriterionScore] = field(default_factory=list)
    overall_percent: Optional[float] = None
    letter: str = "—"
    professor_summary: str = ""
    student_summary: str = ""
    confidence: str = ""
    words: int = 0
    path: str = ""

    def tier_map(self) -> Dict[str, Optional[int]]:
        return {c.name: c.tier for c in self.criteria}


def _parse_json_object(raw: str) -> Optional[dict]:
    """Tolerant JSON extraction — strip code fences and grab the outermost
    object. Same defensive spirit as the participation lab: never trust the
    model to emit clean output, but degrade gracefully."""
    cleaned = re.sub(r"^```(?:json)?|```$", "", raw.strip(), flags=re.MULTILINE).strip()
    start, end = cleaned.find("{"), cleaned.rfind("}")
    if start == -1 or end == -1 or end < start:
        return None
    try:
        return json.loads(cleaned[start:end + 1])
    except json.JSONDecodeError:
        return None


class PaperGrader:
    """Grades each submitted paper against the rubric, grounded in the uploaded
    course material. One grounded LLM call per paper scores every criterion at
    once (so the model reasons over the whole paper coherently); the OVERALL
    score is then computed in code from the tiers and rubric weights — the model
    never does the arithmetic. Missing papers are decided in code, never sent to
    the model."""

    MAX_PAPER_CHARS = 14000            # cap a very long paper's text sent to the model
    RETRIEVAL_QUERY_CHARS = 1500       # cap the retrieval query length

    def __init__(self, retriever, bedrock_runtime=None, rubric: Rubric = None,
                 model_id: str = "amazon.nova-lite-v1:0", llm=None):
        """`llm` is any LangChain chat model (ChatOllama for local runs, or your
        own). If it's None, the grader builds Nova Lite on Bedrock itself — the
        default SageMaker Studio path. `bedrock_runtime` is accepted for API
        symmetry with the participation lab but is not required here."""
        self.retriever = retriever
        self.rubric = rubric
        if llm is None:
            from langchain_aws import ChatBedrockConverse   # lazy: only the Bedrock path needs it
            llm = ChatBedrockConverse(model=model_id, temperature=0.2, max_tokens=3000)
        self._model = llm
        self._chain = PromptTemplate.from_template(GRADING_TEMPLATE) | self._model | StrOutputParser()
        self._rubric_text = "\n\n".join(c.rubric_block() for c in rubric.criteria)

    def _course_context(self, query: str) -> str:
        docs = self.retriever.invoke(query[:self.RETRIEVAL_QUERY_CHARS])
        return "\n\n".join(d.page_content for d in docs)

    def grade_paper(self, display_name: str, sub: Submission) -> PaperGrade:
        """Grade one submitted paper. Retrieves the course material most relevant
        to THIS paper so the 'did they use the concepts correctly' judgment is
        grounded in what was actually taught, not the model's own knowledge."""
        if not sub.text.strip():
            return PaperGrade(
                name=display_name, status="unreadable", confidence="low", path=sub.path,
                professor_summary="No readable text — likely a scanned/image PDF. Grade by hand.",
            )
        context = self._course_context(sub.text)
        raw = self._chain.invoke({
            "course_context": context,
            "rubric_text": self._rubric_text,
            "max_tier": self.rubric.max_tier,
            "student_name": display_name,
            "num_words": sub.words,
            "paper_text": sub.text[:self.MAX_PAPER_CHARS],
        })
        data = _parse_json_object(raw)
        if data is None:
            return PaperGrade(
                name=display_name, status="parse_error", confidence="low",
                words=sub.words, path=sub.path,
                professor_summary="Could not parse the model output — re-run the grading cell.",
            )

        # Keep only criteria that exist in the rubric, in rubric order, so a model
        # that renames/reorders/hallucinates a criterion can't distort the grade.
        by_name = {str(c.get("name", "")).strip().lower(): c for c in data.get("criteria", []) or []}
        scores: List[CriterionScore] = []
        for crit in self.rubric.criteria:
            c = by_name.get(crit.name.strip().lower(), {})
            tier = c.get("tier")
            try:
                tier = int(tier)
                tier = max(0, min(self.rubric.max_tier, tier))
            except (TypeError, ValueError):
                tier = None
            scores.append(CriterionScore(
                name=crit.name,
                tier=tier,
                evidence_quotes=list(c.get("evidence_quotes", []) or []),
                rationale=str(c.get("rationale", "")),
                student_note=str(c.get("student_note", "")),
            ))

        tiers = {s.name: s.tier for s in scores}
        pct = self.rubric.overall_percent(tiers)
        return PaperGrade(
            name=display_name, status="graded", criteria=scores,
            overall_percent=pct, letter=self.rubric.letter(pct),
            professor_summary=str(data.get("professor_summary", "")),
            student_summary=str(data.get("student_summary", "")),
            confidence=str(data.get("confidence", "")),
            words=sub.words, path=sub.path,
        )

    def grade_all(
        self,
        matched: List[Tuple[str, Submission]],
        missing: Optional[List[str]] = None,
    ) -> List[PaperGrade]:
        """Grade every submitted paper, then append missing students as a
        flagged 'missing' row (no score — a non-submission is the professor's
        call, not the model's). Sorted highest overall first."""
        grades = [self.grade_paper(name, sub) for name, sub in matched]
        for name in (missing or []):
            grades.append(PaperGrade(
                name=name, status="missing", confidence="high",
                professor_summary="On the roster but no paper was found. Confirm before entering a 0.",
            ))
        grades.sort(key=lambda g: (g.overall_percent if g.overall_percent is not None else -1),
                    reverse=True)
        return grades


# --------------------------------------------------------------------------- #
# Reporting                                                                   #
# --------------------------------------------------------------------------- #

def _fmt_quotes(quotes: List[str], limit: int = 2) -> str:
    shown = [f'"{q.strip()}"' for q in quotes[:limit] if q and q.strip()]
    return "; ".join(shown) if shown else "—"


def build_report(grades: List[PaperGrade], rubric: Rubric) -> Tuple[str, List[dict]]:
    """Render (markdown_summary, csv_rows). The markdown is a per-student
    summary table plus the per-criterion tiers; the csv_rows feed the take-home
    export (one row per student, one column per criterion)."""
    crit_names = [c.name for c in rubric.criteria]
    header = (
        "| Student | Overall | Grade | "
        + " | ".join(crit_names)
        + " | Confidence |\n|---|:---:|:---:|"
        + "".join(":---:|" for _ in crit_names)
        + ":---:|"
    )
    rows_md: List[str] = []
    csv_rows: List[dict] = []
    for g in grades:
        overall_disp = "—" if g.overall_percent is None else f"{g.overall_percent:g}%"
        grade_disp = g.letter
        if g.status == "missing":
            overall_disp, grade_disp = "—", "missing"
        elif g.status == "unreadable":
            overall_disp, grade_disp = "—", "⚠️ unreadable"
        elif g.status == "parse_error":
            overall_disp, grade_disp = "—", "⚠️ re-run"

        tmap = g.tier_map()
        tier_cells = " | ".join(
            "—" if tmap.get(n) is None else str(tmap[n]) for n in crit_names
        )
        rows_md.append(
            f"| {g.name} | {overall_disp} | {grade_disp} | {tier_cells} | {g.confidence or '—'} |"
        )

        row = {
            "student_name": g.name,
            "status": g.status,
            "overall_percent": "" if g.overall_percent is None else g.overall_percent,
            "letter": g.letter if g.status == "graded" else "",
            "words": g.words,
            "confidence": g.confidence,
            "professor_summary": g.professor_summary,
        }
        for n in crit_names:
            row[f"tier::{n}"] = "" if tmap.get(n) is None else tmap[n]
        csv_rows.append(row)

    return "\n".join([header, *rows_md]), csv_rows


def build_detail(grades: List[PaperGrade], rubric: Rubric) -> str:
    """A longer per-student breakdown for reading in the notebook: the professor
    summary, then each criterion's tier + evidence + rationale, then the note
    drafted for the student."""
    out: List[str] = []
    for g in grades:
        head = f"### {g.name} — "
        if g.status == "graded":
            head += f"{g.overall_percent:g}% ({g.letter}) · confidence {g.confidence or '—'}"
        else:
            head += g.status
        out.append(head)
        if g.professor_summary:
            out.append(f"**For the professor:** {g.professor_summary}")
        for c in g.criteria:
            tier = "—" if c.tier is None else f"{c.tier}/{rubric.max_tier}"
            out.append(
                f"- **{c.name}** — {tier}. {c.rationale} "
                f"_Evidence:_ {_fmt_quotes(c.evidence_quotes)}"
            )
        if g.student_summary:
            out.append(f"**Draft note to student:** {g.student_summary}")
        out.append("")
    return "\n".join(out)


def write_student_feedback(grades: List[PaperGrade], rubric: Rubric, folder: str) -> List[str]:
    """Write one markdown feedback file per graded student into `folder`, ready
    to hand back. Contains ONLY the student-facing material (their summary and
    per-criterion student notes) — never the professor's private rationale.
    Returns the list of paths written."""
    os.makedirs(folder, exist_ok=True)
    written: List[str] = []
    for g in grades:
        if g.status != "graded":
            continue
        safe = re.sub(r"[^\w]+", "_", g.name).strip("_") or "student"
        path = os.path.join(folder, f"{safe}.md")
        lines = [
            f"# Feedback — {g.name}",
            "",
            f"**Grade (draft, pending instructor review): {g.overall_percent:g}% ({g.letter})**",
            "",
            g.student_summary or "",
            "",
            "## By criterion",
        ]
        for c in g.criteria:
            tier = "—" if c.tier is None else f"{c.tier}/{rubric.max_tier}"
            lines.append(f"- **{c.name} ({tier}):** {c.student_note}")
        with open(path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines).rstrip() + "\n")
        written.append(path)
    return written

#!/usr/bin/env python3
"""Prove the deck is still subtractive against the legal-approved baseline.

Legal cleared the words on the 28 approved slides. The standing permission is that copy may be
**removed** from them, never added and never reworded — so the test that matters is not "is this
file unchanged" but "does this file contain only approved strings, and fewer of them."

    python3 verify_deck.py                  # check the .pptx against the approved baseline
    python3 verify_deck.py --pdf out.pdf    # check an exported PDF instead
    python3 verify_deck.py --show           # also print every deletion, slide by slide

It checks the approved slides only. Slides added ahead of them are new material rather than an
edit to approved copy, so LEADING_EXTRA exempts them by count. A QR code stamped on an approved
slide passes because it adds no words — which is exactly why the code carries no caption.

Exit status is 0 when the deck is subtractive and 1 when anything was added, so this can gate a
build. `export_pdf.sh` calls it before it will overwrite the approved PDF.

Three kinds of noise are ignored, none of which is slide copy:

  * `<#>` — PowerPoint's slide-number field code. It lives in the .pptx and renders as the
    actual page number in the PDF, so it reads as an addition in every comparison.
  * `-` bullet glyphs, which the PDF draws and python-pptx does not report.
  * The printed slide number, which the PDF draws as text and the .pptx holds as that field.

One false deletion survives by design: slide 28's "Thank you!" lives in the template layout
rather than in a shape, so python-pptx does not report it and a .pptx check reads it as
removed. A --pdf check sees it correctly. Prefer --pdf for the final pre-delivery verification.

Whitespace is ignored entirely: the PDF text extractor runs adjacent text runs together
("Openyour"), so only the character stream is meaningful.

BASELINE is the approved artifact and is itself frozen. Do not regenerate it from a rebuilt
deck — that would rebase the check onto whatever the deck currently says and prove nothing.
"""

import argparse
import difflib
import re
import sys
from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE_TYPE

HERE = Path(__file__).resolve().parent
BASELINE = HERE / "approved-deck-baseline.txt"
DEFAULT_PPTX = HERE / "assessment-lifecycle-deck.pptx"

# Slides added ahead of the approved 28. NEW material, which legal has not seen and which this
# check deliberately does not police — it exists to protect the approved copy from being edited,
# not to forbid new slides. Currently 1: the shared-folder slide, numbered 0 in the deck.
#
# Raising this number exempts another slide from the check, so it is the one line here worth
# reviewing carefully. Anything added must be a genuinely new slide; moving approved copy onto
# a new slide to dodge the check would defeat the point of having it.
LEADING_EXTRA = 1

# Not slide copy: a field code that renders as the page number, and bullet glyphs the PDF
# draws but python-pptx does not report.
NOISE = re.compile(r"‹#›|[•●▪]")

_PUNCT = {
    "’": "'", "‘": "'", "“": '"', "”": '"',
    "‑": "-", "–": "-", "—": "-", "−": "-", " ": "",
}


def canon(text: str) -> str:
    """Collapse a slide to the character stream that carries its approved meaning."""
    for a, b in _PUNCT.items():
        text = text.replace(a, b)
    return re.sub(r"\s+", "", NOISE.sub("", text))


def from_pdf(path: Path) -> list[str]:
    from pypdf import PdfReader
    return [canon(p.extract_text() or "") for p in PdfReader(str(path)).pages]


def from_pptx(path: Path) -> list[str]:
    from pptx import Presentation

    def shape_text(sh) -> str:
        out = []
        if sh.has_text_frame:
            out.append(sh.text_frame.text)
        if sh.has_table:
            out += [c.text for r in sh.table.rows for c in r.cells]
        return "".join(out)

    return [canon("".join(shape_text(sh) for sh in s.shapes))
            for s in Presentation(str(path)).slides]


def check_qr(pptx: Path) -> int:
    """A separate contract from the legal one: the folder code is on every slide.

    Attendees build all four tools from prompts in the shared folder, so a slide without the
    code is a slide where a latecomer, or anyone who looked up at the wrong moment, has no way
    back to the materials. Cheap to check, and easy to lose when a slide is added by hand in
    PowerPoint rather than through build_deck.py.

    It also decodes the asset, so a stale cached PNG pointing at last year's folder fails here
    rather than in the room. Decoding needs OpenCV; without it the check degrades to presence.
    """
    from pptx import Presentation
    from session_config import DRIVE_URL

    # HERE, not pptx.parent: the asset belongs to the project, and the target may be a copy
    # somewhere else. Resolving it beside the target made a stray .pptx report "asset missing"
    # instead of the failure it actually had.
    qr = HERE / "assets" / "drive-qr.png"
    if not qr.exists():
        print(f"FAIL  {qr.name} is missing; run build_deck.py.", file=sys.stderr)
        return 1

    # Match the image bytes, not merely "has a picture" — most slides also carry the template's
    # book-and-laptop footer mark, so a presence count would pass a slide that lost its code.
    want = qr.read_bytes()
    missing = [n for n, s in enumerate(Presentation(str(pptx)).slides)
               if not any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE
                          and sh.image.blob == want for sh in s.shapes)]
    if missing:
        print(f"FAIL  no shared-folder QR on slide(s): {', '.join(map(str, missing))}",
              file=sys.stderr)
        return 1
    try:
        import cv2
        value, *_ = cv2.QRCodeDetector().detectAndDecode(cv2.imread(str(qr)))
    except ImportError:
        print("OK    QR present on every slide (install opencv-python to also decode it).")
        return 0
    if value != DRIVE_URL:
        print(f"FAIL  {qr.name} decodes to {value or '(nothing)'},\n      not the configured "
              f"DRIVE_URL {DRIVE_URL}", file=sys.stderr)
        return 1
    print(f"OK    QR present on every slide, and decodes to {value}")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--pdf", type=Path, help="check this exported PDF instead of the .pptx")
    ap.add_argument("--pptx", type=Path, default=DEFAULT_PPTX)
    ap.add_argument("--show", action="store_true", help="print every deletion")
    args = ap.parse_args()

    if not BASELINE.exists():
        print(f"No approved baseline at {BASELINE.name} — cannot verify.", file=sys.stderr)
        return 1

    raw = BASELINE.read_text(encoding="utf-8")
    # split("\n"), not splitlines(): splitlines() treats the form feed that separates the
    # slide records as a line break and would destroy them.
    raw = "\n".join(ln for ln in raw.split("\n") if not ln.startswith("#"))
    approved = [canon(s) for s in raw.split("\n\x0c\n") if s.strip()]
    target = args.pdf or args.pptx
    slides = from_pdf(args.pdf) if args.pdf else from_pptx(args.pptx)

    if len(slides) != len(approved) + LEADING_EXTRA:
        print(f"FAIL  {target.name} has {len(slides)} slides; expected {len(approved)} approved "
              f"+ {LEADING_EXTRA} new. Approved slides may not be added or removed.",
              file=sys.stderr)
        return 1
    new_slides, slides = slides[:LEADING_EXTRA], slides[LEADING_EXTRA:]

    additions, deletions = [], []
    for n, (a, b) in enumerate(zip(approved, slides), 1):
        if a == b:
            continue
        for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b,
                                                           autojunk=False).get_opcodes():
            if tag in ("delete", "replace") and a[i1:i2].strip():
                # The PDF draws the slide number as text; the .pptx holds a field code for it,
                # which canon() strips. Every slide therefore "loses" its own number. Not copy.
                if a[i1:i2] != str(n):
                    deletions.append((n, a[i1:i2]))
            if tag in ("insert", "replace") and b[j1:j2].strip():
                additions.append((n, b[j1:j2]))

    if args.show or additions:
        for n, txt in deletions:
            print(f"  slide {n:>2}  - {txt[:120]}")
    for n, txt in additions:
        print(f"  slide {n:>2}  + {txt[:120]}   <== ADDED", file=sys.stderr)

    if additions:
        print(f"\nFAIL  {len(additions)} addition(s) to legal-approved slides in "
              f"{target.name}.\n      Copy may be removed, never added or reworded. "
              f"See DECK-APPROVAL.md.", file=sys.stderr)
        return 1

    if not args.pdf and check_qr(args.pptx) != 0:
        return 1

    print(f"OK    {target.name} is subtractive against the approved deck: "
          f"{len(deletions)} deletion(s), 0 additions."
          + (f"\n      ({len(new_slides)} new slide(s) ahead of the approved 28, not checked "
             f"\u2014 see LEADING_EXTRA.)" if new_slides else ""))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
